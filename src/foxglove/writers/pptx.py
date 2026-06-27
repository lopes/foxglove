from pathlib import Path
from typing import cast

from pptx import Presentation
from pptx.text.text import TextFrame

from foxglove.metadata import FileMetadata


def write_pptx(content: str, path: Path, *, metadata: FileMetadata | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    slides_content = content.split("---")
    for slide_text in slides_content:
        lines = [line for line in slide_text.strip().split("\n") if line.strip()]
        if not lines:
            continue
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title_shape = slide.shapes.title
        if title_shape is not None:
            title_shape.text = lines[0]
        body = slide.placeholders[1]
        tf = cast(TextFrame, getattr(body, "text_frame"))
        tf.clear()
        for i, line in enumerate(lines[1:]):
            if i == 0:
                tf.text = line
            else:
                tf.add_paragraph().text = line
    if metadata is not None:
        props = prs.core_properties
        props.author = metadata.author
        props.last_modified_by = metadata.author
        props.created = metadata.created_at
        props.modified = metadata.modified_at
    prs.save(str(path))
